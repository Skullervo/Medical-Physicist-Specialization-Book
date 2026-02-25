"""Import exam questions from the Excel file into ExamQuestion model."""
import os
from django.core.management.base import BaseCommand
from sisalto.models import ExamQuestion


# Map Excel area names to standardized names
AREA_NAME_MAP = {
    'Anatomia/Fysiologia': 'Anatomia',
    'Anatomia & fysiologia': 'Anatomia',
    'Radiologia': 'Radiologia',
    'Kliininen neurofysiologia': 'KNF',
    'Sädehoito': 'Sädehoito',
    'Sadehoito': 'Sädehoito',
    'Isotooppilääketiede': 'Isotooppilääketiede',
    'Isotooppil��ketiede': 'Isotooppilääketiede',
    'Kliininen fysiologia': 'Kliininen fysiologia',
    'Fysiologia': 'Fysiologia',
}

MONTH_TO_SEMESTER = {
    1: 'kevät', 2: 'kevät', 3: 'kevät', 4: 'kevät', 5: 'touko', 6: 'kevät',
    7: 'syksy', 8: 'syksy', 9: 'syksy', 10: 'syksy', 11: 'syksy', 12: 'syksy',
}


class Command(BaseCommand):
    help = 'Import exam questions from Sairaalafyysikkotentin_kysymykset_järjestetty.xlsx'

    def add_arguments(self, parser):
        parser.add_argument(
            '--clear', action='store_true',
            help='Delete all existing ExamQuestion objects before importing',
        )
        parser.add_argument(
            '--dry-run', action='store_true',
            help='Show what would be imported without saving',
        )
        parser.add_argument(
            '--file', type=str,
            default='Sairaalafyysikkotentin_kysymykset_jarjestetty.xlsx',
            help='Path to the Excel file',
        )

    def handle(self, *args, **options):
        try:
            import openpyxl
        except ImportError:
            self.stderr.write('openpyxl is required: pip install openpyxl')
            return

        filepath = options['file']
        if not os.path.exists(filepath):
            self.stderr.write(f'File not found: {filepath}')
            return

        dry_run = options['dry_run']
        if options['clear'] and not dry_run:
            deleted, _ = ExamQuestion.objects.all().delete()
            self.stdout.write(f'Deleted {deleted} existing ExamQuestion objects.')

        wb = openpyxl.load_workbook(filepath, data_only=True)

        # Use "Aloittain" sheet (area-organized)
        if 'Aloittain' not in wb.sheetnames:
            self.stderr.write('Sheet "Aloittain" not found in workbook')
            return

        ws = wb['Aloittain']
        questions = self._import_aloittain(ws)

        if dry_run:
            self.stdout.write(f'\n=== DRY RUN - Would import {len(questions)} questions ===')
            areas = {}
            years = set()
            for q in questions:
                area = q['subject_area'] or 'Luokittelematon'
                areas[area] = areas.get(area, 0) + 1
                if q['year']:
                    years.add(q['year'])
            if years:
                self.stdout.write(f'Years: {min(years)}-{max(years)}')
            for area, count in sorted(areas.items()):
                self.stdout.write(f'  {area}: {count} questions')
            return

        # Save to database
        created = 0
        for q in questions:
            ExamQuestion.objects.create(**q)
            created += 1

        self.stdout.write(self.style.SUCCESS(
            f'Imported {created} exam questions.'
        ))

    def _import_aloittain(self, ws):
        """
        Import from the "Aloittain" sheet.

        Structure:
        Row 1: Headers (Ala, Vuosi, Kuukausi, Kysymys)
        Row 2+: Data
        """
        questions = []
        question_number_per_exam = {}

        for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            area = row[0]
            year = row[1]
            month = row[2]
            text = row[3]

            # Skip empty rows
            if not text or not str(text).strip():
                continue

            text = str(text).strip()

            # Standardize area name
            area_str = str(area).strip() if area else ''
            subject_area = AREA_NAME_MAP.get(area_str, area_str)

            # Parse year
            year_int = None
            if year:
                try:
                    year_int = int(year)
                except (ValueError, TypeError):
                    pass

            # Parse month for semester
            semester = ''
            if month:
                try:
                    month_int = int(month)
                    semester = MONTH_TO_SEMESTER.get(month_int, '')
                except (ValueError, TypeError):
                    # Month might be string like "touko"
                    semester = str(month).strip()

            # Build exam_date identifier
            exam_date = ''
            if year_int:
                exam_date = str(year_int)
                if month:
                    try:
                        exam_date += f'/{int(month)}'
                    except (ValueError, TypeError):
                        exam_date += f'/{month}'

            # Track question number per exam
            exam_key = exam_date or 'unknown'
            question_number_per_exam[exam_key] = question_number_per_exam.get(exam_key, 0) + 1

            questions.append({
                'question_text': text,
                'year': year_int,
                'semester': semester,
                'subject_area': subject_area,
                'exam_date': exam_date,
                'question_number': question_number_per_exam[exam_key],
            })

        self.stdout.write(f'Found {len(questions)} questions in "Aloittain" sheet')
        return questions
